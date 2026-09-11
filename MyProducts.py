import sys
import sqlite3
from contextlib import contextmanager
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QApplication,
    QFileDialog,
    QGridLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)


DATABASE_PATH = Path(__file__).with_name("products.db")


@contextmanager
def get_connection():
    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row
    try:
        yield connection
    except Exception:
        connection.rollback()
        raise
    else:
        connection.commit()
    finally:
        connection.close()


def create_table():
    with get_connection() as connection:
        connection.execute(
            """
            CREATE TABLE IF NOT EXISTS Products (
                productID INTEGER PRIMARY KEY AUTOINCREMENT,
                productName TEXT NOT NULL,
                productPrice INTEGER NOT NULL CHECK (productPrice >= 0)
            )
            """
        )


def add_product(product_name, product_price):
    with get_connection() as connection:
        cursor = connection.execute(
            "INSERT INTO Products (productName, productPrice) VALUES (?, ?)",
            (product_name, product_price),
        )
        return cursor.lastrowid


def update_product(product_id, product_name, product_price):
    with get_connection() as connection:
        cursor = connection.execute(
            """
            UPDATE Products
            SET productName = ?, productPrice = ?
            WHERE productID = ?
            """,
            (product_name, product_price, product_id),
        )
        return cursor.rowcount > 0


def delete_product(product_id):
    with get_connection() as connection:
        cursor = connection.execute(
            "DELETE FROM Products WHERE productID = ?",
            (product_id,),
        )
        return cursor.rowcount > 0


def get_product(product_id):
    with get_connection() as connection:
        return connection.execute(
            "SELECT productID, productName, productPrice FROM Products WHERE productID = ?",
            (product_id,),
        ).fetchone()


def search_products(product_name=""):
    with get_connection() as connection:
        return connection.execute(
            """
            SELECT productID, productName, productPrice
            FROM Products
            WHERE productName LIKE ?
            ORDER BY productID
            """,
            (f"%{product_name}%",),
        ).fetchall()


def save_products_to_excel(products, file_path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Products"

    headers = ["productID", "productName", "productPrice"]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center")

    for product in products:
        worksheet.append(
            [product["productID"], product["productName"], product["productPrice"]]
        )

    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["B"].width = 28
    worksheet.column_dimensions["C"].width = 18
    worksheet.freeze_panes = "A2"
    workbook.save(file_path)


class ProductWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Products Studio")
        self.resize(900, 680)
        self.setMinimumSize(760, 560)
        self.selected_product_id = None
        self.current_products = []

        self.name_input = QLineEdit()
        self.name_input.setObjectName("fieldInput")
        self.name_input.setPlaceholderText("상품명을 입력하세요")
        self.price_input = QSpinBox()
        self.price_input.setObjectName("fieldInput")
        self.price_input.setRange(0, 2_147_483_647)
        self.price_input.setSuffix(" 원")
        self.search_input = QLineEdit()
        self.search_input.setObjectName("searchInput")
        self.search_input.setPlaceholderText("상품명 검색")
        self.search_input.textChanged.connect(self.refresh_table)

        self.setStyleSheet("""
            QMainWindow, QWidget#centralWidget { background: #f4f7fb; color: #14213d; font-family: "Segoe UI", "Malgun Gothic"; font-size: 13px; }
            QWidget#headerPanel { background: #14213d; border-radius: 18px; }
            QLabel#eyebrow { color: #55d6be; font-size: 11px; font-weight: 700; }
            QLabel#pageTitle { color: #ffffff; font-size: 28px; font-weight: 800; }
            QLabel#pageSubtitle, QLabel#sectionHint { color: #9fb0c8; font-size: 12px; }
            QLabel#sectionTitle { color: #14213d; font-size: 17px; font-weight: 800; }
            QLabel#countBadge { background: #dff8f2; color: #167c6b; border-radius: 10px; padding: 5px 10px; font-weight: 700; }
            QWidget#contentCard { background: #ffffff; border: 1px solid #e3eaf3; border-radius: 16px; }
            QLabel#fieldLabel { color: #64748b; font-size: 11px; font-weight: 700; }
            QLineEdit#fieldInput, QSpinBox#fieldInput, QLineEdit#searchInput { background: #f8fafc; border: 1px solid #d8e1ec; border-radius: 9px; padding: 10px 12px; color: #14213d; }
            QLineEdit#fieldInput:focus, QSpinBox#fieldInput:focus, QLineEdit#searchInput:focus { border: 2px solid #55d6be; background: #ffffff; }
            QPushButton { min-height: 38px; border: 0; border-radius: 9px; padding: 0 17px; font-weight: 700; }
            QPushButton#primaryButton { background: #55d6be; color: #102a43; }
            QPushButton#primaryButton:hover { background: #73e3ce; }
            QPushButton#secondaryButton { background: #e8eef6; color: #334e68; }
            QPushButton#secondaryButton:hover { background: #d9e3ef; }
            QPushButton#dangerButton { background: #fff0f0; color: #c24141; }
            QPushButton#dangerButton:hover { background: #ffe0e0; }
            QPushButton#exportButton { background: #14213d; color: #ffffff; min-height: 44px; }
            QPushButton#exportButton:hover { background: #21365d; }
            QTableWidget { background: #ffffff; alternate-background-color: #f8fafc; border: 0; gridline-color: #edf1f6; color: #243b53; selection-background-color: #dff8f2; selection-color: #102a43; }
            QHeaderView::section { background: #f4f7fb; border: 0; border-bottom: 1px solid #e3eaf3; color: #8291a5; font-size: 11px; font-weight: 700; padding: 12px 10px; }
            QScrollBar:vertical { background: #f4f7fb; width: 10px; margin: 2px; }
            QScrollBar::handle:vertical { background: #c8d4e3; border-radius: 5px; min-height: 30px; }
        """)

        header_panel = QWidget()
        header_panel.setObjectName("headerPanel")
        header_layout = QVBoxLayout(header_panel)
        header_layout.setContentsMargins(26, 22, 26, 22)
        header_layout.setSpacing(4)
        for text, object_name in (("INVENTORY / DESKTOP", "eyebrow"), ("Products Studio", "pageTitle"), ("제품 데이터를 빠르게 입력하고 한눈에 관리하세요", "pageSubtitle")):
            label = QLabel(text)
            label.setObjectName(object_name)
            header_layout.addWidget(label)

        form_card = QWidget()
        form_card.setObjectName("contentCard")
        form_layout = QVBoxLayout(form_card)
        form_layout.setContentsMargins(20, 18, 20, 18)
        form_layout.setSpacing(14)
        form_heading = QLabel("상품 정보")
        form_heading.setObjectName("sectionTitle")
        form_hint = QLabel("새 상품을 등록하거나 목록에서 선택해 수정하세요")
        form_hint.setObjectName("sectionHint")
        form_layout.addWidget(form_heading)
        form_layout.addWidget(form_hint)

        fields_layout = QGridLayout()
        fields_layout.setHorizontalSpacing(12)
        for column, text in enumerate(("상품명", "상품 가격")):
            label = QLabel(text)
            label.setObjectName("fieldLabel")
            fields_layout.addWidget(label, 0, column)
        fields_layout.addWidget(self.name_input, 1, 0)
        fields_layout.addWidget(self.price_input, 1, 1)
        fields_layout.setColumnStretch(0, 3)
        fields_layout.setColumnStretch(1, 2)
        form_layout.addLayout(fields_layout)

        add_button = QPushButton("상품 입력")
        add_button.setObjectName("primaryButton")
        add_button.clicked.connect(self.add_product_from_form)
        update_button = QPushButton("선택 수정")
        update_button.setObjectName("secondaryButton")
        update_button.clicked.connect(self.update_selected_product)
        delete_button = QPushButton("선택 삭제")
        delete_button.setObjectName("dangerButton")
        delete_button.clicked.connect(self.delete_selected_product)
        clear_button = QPushButton("입력 초기화")
        clear_button.setObjectName("secondaryButton")
        clear_button.clicked.connect(self.clear_form)
        action_layout = QHBoxLayout()
        action_layout.setSpacing(8)
        for button in (add_button, update_button, delete_button, clear_button):
            action_layout.addWidget(button)
        form_layout.addLayout(action_layout)

        list_card = QWidget()
        list_card.setObjectName("contentCard")
        list_layout = QVBoxLayout(list_card)
        list_layout.setContentsMargins(20, 18, 20, 18)
        list_layout.setSpacing(12)
        list_header = QHBoxLayout()
        list_title = QLabel("상품 목록")
        list_title.setObjectName("sectionTitle")
        self.count_badge = QLabel("0 items")
        self.count_badge.setObjectName("countBadge")
        list_header.addWidget(list_title)
        list_header.addStretch()
        list_header.addWidget(self.count_badge)
        list_layout.addLayout(list_header)
        list_layout.addWidget(self.search_input)

        self.product_table = QTableWidget(0, 3)
        self.product_table.setHorizontalHeaderLabels(["productID", "상품명", "상품 가격"])
        self.product_table.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.product_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.product_table.itemSelectionChanged.connect(self.load_selected_product)
        self.product_table.horizontalHeader().setStretchLastSection(True)
        self.product_table.horizontalHeader().setDefaultSectionSize(140)
        self.product_table.verticalHeader().setDefaultSectionSize(42)
        self.product_table.setShowGrid(False)
        self.product_table.setAlternatingRowColors(True)
        list_layout.addWidget(self.product_table)

        export_button = QPushButton("현재 목록 Excel 저장")
        export_button.setObjectName("exportButton")
        export_button.clicked.connect(self.export_to_excel)

        layout = QVBoxLayout()
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(16)
        layout.addWidget(header_panel)
        layout.addWidget(form_card)
        layout.addWidget(list_card, 1)
        layout.addWidget(export_button)
        central_widget = QWidget()
        central_widget.setObjectName("centralWidget")
        central_widget.setLayout(layout)
        self.setCentralWidget(central_widget)
        self.refresh_table()

    def add_product_from_form(self):
        name = self.name_input.text().strip()
        if not name:
            QMessageBox.warning(self, "입력 오류", "상품명을 입력하세요.")
            return

        product_id = add_product(name, self.price_input.value())
        self.clear_form()
        self.refresh_table()
        QMessageBox.information(self, "완료", f"상품이 입력되었습니다. ID: {product_id}")

    def update_selected_product(self):
        if self.selected_product_id is None:
            QMessageBox.warning(self, "선택 오류", "수정할 상품을 선택하세요.")
            return

        name = self.name_input.text().strip()
        if not name:
            QMessageBox.warning(self, "입력 오류", "상품명을 입력하세요.")
            return

        update_product(self.selected_product_id, name, self.price_input.value())
        self.clear_form()
        self.refresh_table()
        QMessageBox.information(self, "완료", "상품이 수정되었습니다.")

    def delete_selected_product(self):
        if self.selected_product_id is None:
            QMessageBox.warning(self, "선택 오류", "삭제할 상품을 선택하세요.")
            return

        answer = QMessageBox.question(
            self,
            "삭제 확인",
            "선택한 상품을 삭제하시겠습니까?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
        )
        if answer == QMessageBox.StandardButton.Yes:
            delete_product(self.selected_product_id)
            self.clear_form()
            self.refresh_table()

    def load_selected_product(self):
        selected_items = self.product_table.selectedItems()
        if not selected_items:
            return

        row = selected_items[0].row()
        self.selected_product_id = int(self.product_table.item(row, 0).text())
        self.name_input.setText(self.product_table.item(row, 1).text())
        self.price_input.setValue(
            int(self.product_table.item(row, 2).data(Qt.ItemDataRole.UserRole))
        )

    def clear_form(self):
        self.selected_product_id = None
        self.name_input.clear()
        self.price_input.setValue(0)
        self.product_table.clearSelection()

    def refresh_table(self):
        self.current_products = search_products(self.search_input.text().strip())
        self.count_badge.setText(f"{len(self.current_products)} items")
        self.product_table.setRowCount(0)
        for product in self.current_products:
            row = self.product_table.rowCount()
            self.product_table.insertRow(row)
            self.product_table.setItem(row, 0, QTableWidgetItem(str(product["productID"])))
            self.product_table.setItem(row, 1, QTableWidgetItem(product["productName"]))
            price_item = QTableWidgetItem(f"{product['productPrice']:,} 원")
            price_item.setData(Qt.ItemDataRole.UserRole, product["productPrice"])
            self.product_table.setItem(row, 2, price_item)

    def export_to_excel(self):
        if not self.current_products:
            QMessageBox.warning(self, "저장할 데이터 없음", "현재 목록이 비어 있습니다.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Excel 파일 저장",
            str(Path.home() / "products.xlsx"),
            "Excel 파일 (*.xlsx)",
        )
        if not file_path:
            return

        try:
            save_products_to_excel(self.current_products, file_path)
            QMessageBox.information(self, "저장 완료", f"Excel 파일을 저장했습니다.\n{file_path}")
        except OSError as error:
            QMessageBox.critical(self, "저장 오류", f"파일을 저장할 수 없습니다.\n{error}")


def run_gui():
    create_table()
    application = QApplication(sys.argv)
    window = ProductWindow()
    window.show()
    sys.exit(application.exec())


if __name__ == "__main__":
    run_gui()