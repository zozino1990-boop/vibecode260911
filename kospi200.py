import argparse
import re
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
import requests
from bs4 import BeautifulSoup
from PyQt6.QtCore import QObject, QThread, pyqtSignal
from PyQt6.QtWidgets import (
    QApplication,
    QAbstractItemView,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMainWindow,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)


INDEX_URL = "https://finance.naver.com/sise/sise_index.naver?code=KPI200"
CONSTITUENTS_URL = "https://finance.naver.com/sise/entryJongmok.naver?type=KPI200"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/140.0.0.0 Safari/537.36"
    )
}
HEADER_NAMES = {
    "종목별": "종목명",
    "거래대금(백만)": "거래대금(백만)",
    "시가총액(억)": "시가총액(억)",
}


def clean_text(element) -> str:
    """HTML 요소 안의 연속된 공백과 줄바꿈을 정리합니다."""
    return re.sub(r"\s+", " ", element.get_text(" ", strip=True)).strip()


def get_header_name(header) -> str:
    name = re.sub(r"\s*\(\s*", "(", clean_text(header))
    name = re.sub(r"\s*\)", ")", name)
    return HEADER_NAMES.get(name, name)


def parse_constituents_page(
    soup: BeautifulSoup, limit: int | None = None
) -> list[dict[str, Any]]:
    """편입종목 한 페이지의 표를 파싱합니다."""
    title = soup.select_one("div.box_type_m h4.top_tlt em")
    box = title.find_parent("div", class_="box_type_m") if title else None
    table = box.select_one("table.type_1") if box else None
    if table is None:
        raise RuntimeError("KPI200 편입종목 표를 찾지 못했습니다.")

    header_row = table.select_one("tr")
    headers = (
        [get_header_name(header) for header in header_row.select("th")]
        if header_row
        else []
    )
    if len(headers) != 7:
        raise RuntimeError("편입종목 표의 헤더 구조가 변경되었습니다.")

    constituents: list[dict[str, Any]] = []
    for row in table.select("tr"):
        cells = row.select("td")
        if len(cells) != len(headers):
            continue

        name_link = cells[0].select_one("a")
        name = cells[0].get_text(" ", strip=True)
        if not name:
            continue

        item = {column: clean_text(cell) for column, cell in zip(headers, cells)}
        item["종목명"] = name
        change_direction = cells[2].select_one("span.blind")
        change_value_element = cells[2].select_one("span.tah")
        change_value = (
            clean_text(change_value_element)
            if change_value_element
            else clean_text(cells[2])
        )
        if change_direction:
            item["전일비"] = f"{clean_text(change_direction)} {change_value}"
        item["종목코드"] = (
            re.search(r"code=([0-9]+)", name_link["href"]).group(1)
            if name_link and name_link.has_attr("href")
            and re.search(r"code=([0-9]+)", name_link["href"])
            else ""
        )
        item["상세페이지"] = (
            urljoin("https://finance.naver.com", name_link["href"])
            if name_link and name_link.has_attr("href")
            else ""
        )
        constituents.append(item)

        if limit is not None and len(constituents) >= limit:
            break

    return constituents


def fetch_constituents(limit: int | None = 200) -> list[dict[str, Any]]:
    """네이버 금융 KPI200 편입종목을 여러 페이지에서 가져옵니다."""
    constituents: list[dict[str, Any]] = []

    with requests.Session() as session:
        page = 1
        while limit is None or len(constituents) < limit:
            response = session.get(
                CONSTITUENTS_URL,
                params={"page": page},
                headers={**HEADERS, "Referer": INDEX_URL},
                timeout=15,
            )
            response.raise_for_status()

            soup = BeautifulSoup(
                response.content,
                "html.parser",
                from_encoding=response.apparent_encoding,
            )
            page_items = parse_constituents_page(
                soup,
                limit=None if limit is None else limit - len(constituents),
            )
            if not page_items:
                break

            constituents.extend(page_items)
            page += 1

    return constituents


def save_constituents_to_xlsx(
    constituents: list[dict[str, Any]], file_path: str = "kospi200.xlsx"
) -> Path:
    """편입종목 데이터를 Excel 파일로 저장합니다."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "KOSPI200"
    columns = [
        "번호",
        "종목코드",
        "종목명",
        "현재가",
        "전일비",
        "등락률",
        "거래량",
        "거래대금(백만)",
        "시가총액(억)",
    ]

    worksheet.append(columns)
    for row_number, item in enumerate(constituents, start=1):
        worksheet.append(
            [
                row_number,
                item["종목코드"],
                item["종목명"],
                item["현재가"],
                item["전일비"],
                item["등락률"],
                item["거래량"],
                item["거래대금(백만)"],
                item["시가총액(억)"],
            ]
        )

    header_fill = PatternFill("solid", fgColor="12304A")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    widths = [8, 12, 18, 14, 14, 12, 16, 18, 16]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + column)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    output_path = Path(file_path)
    workbook.save(output_path)
    return output_path.resolve()


class CrawlWorker(QObject):
    finished = pyqtSignal(list)
    failed = pyqtSignal(str)

    def __init__(self, limit: int | None):
        super().__init__()
        self.limit = limit

    def run(self) -> None:
        try:
            self.finished.emit(fetch_constituents(self.limit))
        except Exception as error:
            self.failed.emit(str(error))


class NumericTableWidgetItem(QTableWidgetItem):
    """쉼표와 상승·하락 문자가 포함된 값을 숫자로 정렬합니다."""

    def __init__(self, text: str):
        super().__init__(text)
        self.sort_value = self.get_sort_value(text)

    @staticmethod
    def get_sort_value(text: str) -> float:
        number_match = re.search(r"[-+]?\d+(?:\.\d+)?", text.replace(",", ""))
        if not number_match:
            return float("-inf")

        value = float(number_match.group())
        if "하락" in text:
            return -abs(value)
        if "상승" in text:
            return abs(value)
        return value

    def __lt__(self, other: QTableWidgetItem) -> bool:
        if isinstance(other, NumericTableWidgetItem):
            return self.sort_value < other.sort_value
        return super().__lt__(other)


class Kospi200Window(QMainWindow):
    columns = [
        "번호",
        "종목코드",
        "종목명",
        "현재가",
        "전일비",
        "등락률",
        "거래량",
        "거래대금(백만)",
        "시가총액(억)",
    ]

    def __init__(self, limit: int | None = 200):
        super().__init__()
        self.limit = limit
        self.constituents: list[dict[str, Any]] = []
        self.thread: QThread | None = None
        self.worker: CrawlWorker | None = None
        self.setWindowTitle("KOSPI 200 편입종목상위")
        self.resize(1440, 820)
        self.build_ui()
        self.load_constituents()

    def build_ui(self) -> None:
        central_widget = QWidget()
        layout = QVBoxLayout(central_widget)
        layout.setContentsMargins(28, 24, 28, 24)
        layout.setSpacing(14)

        self.setStyleSheet(
            """
            QMainWindow, QWidget { background: #F3F7FA; color: #183247; }
            QLabel#title { color: #12304A; font-size: 28px; font-weight: 800; }
            QLabel#subtitle { color: #6A7D8D; font-size: 13px; }
            QLabel#status { color: #168A78; font-weight: 700; padding: 8px 0; }
            QLineEdit { background: #FFFFFF; border: 1px solid #CAD8E1;
                border-radius: 8px; padding: 10px 12px; }
            QLineEdit:focus { border: 2px solid #1AA58A; }
            QPushButton { background: #12304A; color: #FFFFFF; border: 0;
                border-radius: 8px; padding: 10px 18px; font-weight: 700; }
            QPushButton:hover { background: #1AA58A; }
            QPushButton:disabled { background: #AABBC5; }
            QTableWidget { background: #FFFFFF; alternate-background-color: #F6FAFC;
                border: 1px solid #D5E1E8; border-radius: 10px;
                gridline-color: #E8EFF3; selection-background-color: #CDEFE7;
                selection-color: #12304A; }
            QHeaderView::section { background: #DCEAF0; color: #315368;
                border: 0; padding: 10px 8px; font-weight: 800; }
            """
        )

        title = QLabel("KOSPI 200 편입종목상위")
        title.setObjectName("title")
        subtitle = QLabel("네이버 금융 실시간 표를 수집해 한눈에 확인하고 Excel로 저장합니다.")
        subtitle.setObjectName("subtitle")

        toolbar = QHBoxLayout()
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("종목명 또는 종목코드 검색")
        self.search_input.textChanged.connect(self.filter_table)
        self.refresh_button = QPushButton("새로고침")
        self.refresh_button.setToolTip("KOSPI200 데이터를 다시 가져옵니다.")
        self.refresh_button.clicked.connect(self.load_constituents)
        self.save_button = QPushButton("Excel 저장")
        self.save_button.setToolTip("현재 목록을 kospi200.xlsx로 저장합니다.")
        self.save_button.setEnabled(False)
        self.save_button.clicked.connect(self.save_excel)
        toolbar.addWidget(self.search_input, 1)
        toolbar.addWidget(self.refresh_button)
        toolbar.addWidget(self.save_button)

        self.status_label = QLabel("데이터를 불러오는 중입니다...")
        self.status_label.setObjectName("status")

        self.table = QTableWidget(0, len(self.columns))
        self.table.setHorizontalHeaderLabels(self.columns)
        self.table.setSortingEnabled(True)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setSelectionBehavior(
            QAbstractItemView.SelectionBehavior.SelectRows
        )
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        self.table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.ResizeToContents
        )
        self.table.horizontalHeader().setStretchLastSection(True)

        layout.addWidget(title)
        layout.addWidget(subtitle)
        layout.addLayout(toolbar)
        layout.addWidget(self.status_label)
        layout.addWidget(self.table)
        self.setCentralWidget(central_widget)

    def load_constituents(self) -> None:
        if self.thread and self.thread.isRunning():
            return

        self.refresh_button.setEnabled(False)
        self.save_button.setEnabled(False)
        self.status_label.setText("KOSPI200 종목을 불러오는 중입니다...")
        self.constituents = []
        self.table.setRowCount(0)

        self.thread = QThread(self)
        self.worker = CrawlWorker(self.limit)
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.run)
        self.worker.finished.connect(self.display_constituents)
        self.worker.failed.connect(self.show_error)
        self.worker.finished.connect(self.thread.quit)
        self.worker.failed.connect(self.thread.quit)
        self.worker.finished.connect(self.worker.deleteLater)
        self.worker.failed.connect(self.worker.deleteLater)
        self.thread.finished.connect(self.thread.deleteLater)
        self.thread.finished.connect(self.loading_finished)
        self.thread.start()

    def display_constituents(self, constituents: list[dict[str, Any]]) -> None:
        self.constituents = constituents
        self.table.setSortingEnabled(False)
        self.table.setRowCount(len(constituents))
        for row_number, item in enumerate(constituents):
            values = [
                str(row_number + 1),
                item["종목코드"],
                item["종목명"],
                item["현재가"],
                item["전일비"],
                item["등락률"],
                item["거래량"],
                item["거래대금(백만)"],
                item["시가총액(억)"],
            ]
            for column_number, value in enumerate(values):
                item_class = (
                    NumericTableWidgetItem
                    if column_number == 0 or 3 <= column_number <= 8
                    else QTableWidgetItem
                )
                self.table.setItem(
                    row_number,
                    column_number,
                    item_class(value),
                )
        self.table.setSortingEnabled(True)
        self.status_label.setText(f"총 {len(constituents)}개 종목")
        self.save_button.setEnabled(bool(constituents))
        self.filter_table(self.search_input.text())

    def filter_table(self, text: str) -> None:
        keyword = text.strip().lower()
        for row in range(self.table.rowCount()):
            name = self.table.item(row, 2).text().lower()
            code = self.table.item(row, 1).text().lower()
            self.table.setRowHidden(row, keyword not in name and keyword not in code)

    def save_excel(self) -> None:
        if not self.constituents:
            return
        try:
            output_path = save_constituents_to_xlsx(self.constituents)
            self.status_label.setText(f"저장 완료: {output_path.name}")
        except Exception as error:
            self.status_label.setText(f"Excel 저장 실패: {error}")

    def show_error(self, message: str) -> None:
        self.status_label.setText(f"크롤링 실패: {message}")

    def loading_finished(self) -> None:
        self.refresh_button.setEnabled(True)
        self.thread = None
        self.worker = None


def main() -> None:
    parser = argparse.ArgumentParser(
        description="네이버 금융 KPI200 편입종목상위 정보를 크롤링합니다."
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=200,
        help="출력할 종목 수입니다. 기본값은 200개이며, 0은 전체 페이지입니다.",
    )
    args = parser.parse_args()

    if args.limit < 0:
        parser.error("--limit은 0 이상이어야 합니다.")

    limit = None if args.limit == 0 else args.limit
    application = QApplication([])
    window = Kospi200Window(limit=limit)
    window.show()
    application.exec()


if __name__ == "__main__":
    main()
